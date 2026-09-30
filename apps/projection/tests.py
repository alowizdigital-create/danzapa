import io
import shutil
import tempfile
from datetime import date

from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from PIL import Image
from pptx import Presentation

from apps.chants.models import Chant
from apps.comptes import roles
from apps.comptes.models import Utilisateur
from apps.cultes import diapos as d
from apps.cultes.models import Culte, ElementCulte

from . import pptx as export
from . import rendu
from .models import CLASSIQUE, Theme

SANS_ATTENDRE = """1. Sans attendre
Je veux tendre
Au bonheur promis

Refrain
De mon Dieu je suis l'enfant
Et c'est lui qui me défend

2. Qui s'élance
Qui s'avance
Obtiendra le prix
"""

MEDIA_TEMP = tempfile.mkdtemp()


def image_png(couleur=(200, 30, 30)):
    tampon = io.BytesIO()
    Image.new("RGB", (32, 18), couleur).save(tampon, "PNG")
    return SimpleUploadedFile("fond.png", tampon.getvalue(), content_type="image/png")


def utilisateur(nom, role):
    u = Utilisateur.objects.create_user(nom, password="x")
    u.groups.add(Group.objects.get(name=role))
    return u


def textes(diapositive):
    return [
        p.text
        for forme in diapositive.shapes
        if forme.has_text_frame
        for p in forme.text_frame.paragraphs
        if p.text
    ]


def ouvrir(contenu):
    return Presentation(io.BytesIO(contenu))


class ThemeParDefautTests(TestCase):
    def test_classique_cree_par_migration(self):
        theme = Theme.objects.get(nom=CLASSIQUE)
        self.assertEqual((theme.couleur_fond, theme.couleur_texte, theme.gras), ("#000000", "#ffffff", True))

    def test_culte_sans_theme_utilise_classique(self):
        culte = Culte.objects.create(date=date(2026, 9, 20))
        self.assertEqual(culte.theme_effectif.nom, CLASSIQUE)


class TailleTests(TestCase):
    def setUp(self):
        self.theme = Theme.par_defaut()

    def test_taille_du_theme_si_le_texte_tient(self):
        diapo = d.Diapo(d.PAROLES, ["Sans attendre", "Je veux tendre"])
        self.assertEqual(rendu.taille_pt(diapo, self.theme), 54)

    def test_ligne_longue_reduite(self):
        diapo = d.Diapo(d.PAROLES, ["Une très longue ligne de paroles qui ne tiendrait jamais sur la diapo"])
        taille = rendu.taille_pt(diapo, self.theme)
        self.assertLess(taille, 54)
        self.assertGreaterEqual(taille, rendu.TAILLE_MIN)

    def test_trop_de_lignes_reduit(self):
        diapo = d.Diapo(d.PAROLES, ["court"] * 10)
        self.assertLess(rendu.taille_pt(diapo, self.theme), 54)

    def test_minimum_sans_depasser_la_taille_du_theme(self):
        self.theme.taille_paroles = 20
        diapo = d.Diapo(d.PAROLES, ["court"])
        self.assertEqual(rendu.taille_pt(diapo, self.theme), 20)

    def test_conversion_cqw(self):
        self.assertEqual(rendu.en_cqw(96), "10.000cqw")


@override_settings(MEDIA_ROOT=MEDIA_TEMP)
class ExportTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_TEMP, ignore_errors=True)

    def setUp(self):
        self.chant = Chant.objects.create(titre="Sans attendre je veux tendre")
        self.chant.remplacer_paroles(SANS_ATTENDRE)
        self.culte = Culte.objects.create(date=date(2026, 9, 20))
        self.element = ElementCulte.objects.create(
            culte=self.culte, ordre=1, type=ElementCulte.CHANT, chant=self.chant, moment="Cantique d'entrée"
        )
        ElementCulte.objects.create(
            culte=self.culte, ordre=2, type=ElementCulte.TEXTE, moment="Annonces", contenu="Réunion mardi", masque=True
        )

    def test_une_diapo_par_diapo_de_l_editeur(self):
        presentation = ouvrir(export.exporter(self.culte))
        attendues = d.diapos_du_culte(self.culte)
        self.assertEqual(len(presentation.slides), len(attendues))
        self.assertEqual((presentation.slide_width, presentation.slide_height), (12192000, 6858000))

    def test_textes_majuscules_et_numeros(self):
        diapositives = list(ouvrir(export.exporter(self.culte)).slides)
        self.assertEqual(textes(diapositives[0]), ["BIENVENUS", "DANS LA", "MAISON DU SEIGNEUR"])
        self.assertEqual(textes(diapositives[1]), ["CANTIQUE D'ENTRÉE : SANS ATTENDRE JE VEUX TENDRE"])
        self.assertEqual(textes(diapositives[2]), ["1. Sans attendre", "Je veux tendre", "Au bonheur promis"])

    def test_mise_en_forme_du_theme(self):
        theme = Theme.objects.create(nom="Bleu", couleur_fond="#102040", couleur_texte="#ffee00", police="Arial", gras=False)
        self.culte.theme = theme
        diapositive = list(ouvrir(export.exporter(self.culte)).slides)[2]
        police = [f for f in diapositive.shapes if f.has_text_frame][0].text_frame.paragraphs[0].runs[0].font
        self.assertEqual((police.name, police.bold, str(police.color.rgb)), ("Arial", False, "FFEE00"))
        self.assertEqual(str(diapositive.background.fill.fore_color.rgb), "102040")

    def test_elements_masques_deviennent_diapos_masquees(self):
        diapositives = list(ouvrir(export.exporter(self.culte)).slides)
        masquees = [s for s in diapositives if s._element.get("show") == "0"]
        self.assertEqual([textes(s) for s in masquees], [["ANNONCES"], ["Réunion mardi"]])

    def test_notes_du_presentateur(self):
        diapositive = list(ouvrir(export.exporter(self.culte)).slides)[2]
        self.assertEqual(diapositive.notes_slide.notes_text_frame.text, "Sans attendre je veux tendre — Couplet 1")

    def test_images_de_fond_et_de_bienvenue(self):
        theme = Theme.objects.create(nom="Photo", image_fond=image_png(), image_bienvenue=image_png((0, 0, 200)))
        self.culte.theme = theme
        diapositives = list(ouvrir(export.exporter(self.culte)).slides)
        images = lambda s: [f for f in s.shapes if f.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        self.assertEqual(len(images(diapositives[0])), 2)  # fond + bienvenue, sans texte
        self.assertEqual(textes(diapositives[0]), [])
        self.assertEqual(len(images(diapositives[2])), 1)

    def test_nom_de_fichier(self):
        self.culte.titre = "Culte du 20/09 : Sainte Cène"
        self.assertEqual(export.nom_de_fichier(self.culte), "Culte du 20-09 - Sainte Cène.pptx")


@override_settings(MEDIA_ROOT=MEDIA_TEMP)
class VuesTests(TestCase):
    def setUp(self):
        self.editeur = utilisateur("editeur", roles.EDITEUR)
        self.lecteur = utilisateur("lecteur", roles.LECTEUR)
        self.culte = Culte.objects.create(date=date(2026, 9, 20))
        self.export = reverse("cultes:export_pptx", args=[self.culte.pk])

    def test_lecteur_peut_exporter(self):
        self.client.force_login(self.lecteur)
        reponse = self.client.get(self.export)
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(
            reponse["Content-Type"], "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        self.assertIn("attachment", reponse["Content-Disposition"])
        self.assertIn('filename="Culte du 20 septembre 2026.pptx"', reponse["Content-Disposition"])
        self.assertEqual(len(ouvrir(reponse.content).slides), 1)

    def test_nom_accentue_telecharge_en_utf8(self):
        self.culte.titre = "Culte de Pâques"
        self.culte.save()
        self.client.force_login(self.lecteur)
        reponse = self.client.get(self.export)
        self.assertIn("filename*=utf-8''Culte%20de%20P%C3%A2ques.pptx", reponse["Content-Disposition"])

    def test_anonyme_redirige(self):
        self.assertEqual(self.client.get(self.export).status_code, 302)

    def test_choix_du_theme_dans_l_editeur(self):
        theme = Theme.objects.create(nom="Bleu", couleur_fond="#102040")
        self.client.force_login(self.editeur)
        reponse = self.client.post(
            reverse("cultes:parametres", args=[self.culte.pk]), {"theme": theme.pk}, headers={"HX-Request": "true"}
        )
        self.assertContains(reponse, "--fond-couleur: #102040")
        self.culte.refresh_from_db()
        self.assertEqual(self.culte.theme, theme)

    def test_editeur_cree_et_modifie_un_theme(self):
        self.client.force_login(self.editeur)
        donnees = {
            "nom": "Nuit", "couleur_fond": "#001122", "couleur_texte": "#ffffff", "police": "Verdana",
            "gras": "on", "taille_paroles": "48", "police_titres": "Georgia", "taille_titre": "30",
            "couleur_bandeau": "#555555", "image_fond": image_png(),
        }
        reponse = self.client.post(reverse("projection:theme_creation"), donnees)
        self.assertRedirects(reponse, reverse("projection:theme_liste"))
        theme = Theme.objects.get(nom="Nuit")
        self.assertTrue(theme.image_fond.name.startswith("themes/"))
        self.assertFalse(theme.titres_majuscules)

    def test_couleur_invalide_refusee(self):
        self.client.force_login(self.editeur)
        reponse = self.client.post(
            reverse("projection:theme_creation"),
            {"nom": "X", "couleur_fond": "rouge", "couleur_texte": "#ffffff", "police": "Arial",
             "taille_paroles": "48", "police_titres": "Arial", "taille_titre": "30", "couleur_bandeau": "#555555"},
        )
        self.assertContains(reponse, "#RRGGBB")

    def test_lecteur_ne_modifie_pas_les_themes(self):
        self.client.force_login(self.lecteur)
        self.assertEqual(self.client.get(reverse("projection:theme_liste")).status_code, 200)
        self.assertEqual(self.client.get(reverse("projection:theme_creation")).status_code, 403)

    def test_classique_ne_peut_pas_etre_supprime_ni_renomme(self):
        self.client.force_login(self.editeur)
        classique = Theme.par_defaut()
        self.client.post(reverse("projection:theme_suppression", args=[classique.pk]))
        self.assertTrue(Theme.objects.filter(pk=classique.pk).exists())
        reponse = self.client.post(
            reverse("projection:theme_modification", args=[classique.pk]),
            {"nom": "Autre", "couleur_fond": "#000000", "couleur_texte": "#ffffff", "police": "Calibri",
             "taille_paroles": "54", "police_titres": "Georgia", "taille_titre": "32", "couleur_bandeau": "#8f8f8f"},
        )
        self.assertContains(reponse, "ne peut pas être renommé")

    def test_suppression_repasse_les_cultes_au_classique(self):
        theme = Theme.objects.create(nom="Bleu")
        self.culte.theme = theme
        self.culte.save()
        self.client.force_login(self.editeur)
        self.client.post(reverse("projection:theme_suppression", args=[theme.pk]))
        self.culte.refresh_from_db()
        self.assertIsNone(self.culte.theme)
        self.assertEqual(self.culte.theme_effectif.nom, CLASSIQUE)

    def test_duplication_garde_le_theme(self):
        theme = Theme.objects.create(nom="Bleu")
        self.culte.theme = theme
        self.culte.save()
        copie = self.culte.dupliquer(date(2026, 9, 27))
        self.assertEqual(copie.theme, theme)
