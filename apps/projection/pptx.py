"""Export d'un culte en présentation PowerPoint (.pptx).

Les diapos sont les mêmes que dans l'éditeur (`apps.cultes.diapos`) et les
tailles de texte les mêmes que dans l'aperçu (`rendu.taille_pt`).
"""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from apps.cultes import diapos as d

from . import rendu
from .models import Theme

EMU_PAR_PT = 12700
LARGEUR = Emu(rendu.LARGEUR_PT * EMU_PAR_PT)  # 13,33 pouces
HAUTEUR = Emu(rendu.HAUTEUR_PT * EMU_PAR_PT)  # 7,5 pouces
MISE_EN_PAGE_VIERGE = 6


def rgb(hexa):
    return RGBColor.from_string(hexa.lstrip("#").upper())


def pt(valeur):
    return Emu(int(valeur * EMU_PAR_PT))


def lire_image(champ):
    if not champ:
        return None
    with champ.open("rb") as fichier:
        return fichier.read()


class Exporteur:
    def __init__(self, culte):
        self.culte = culte
        self.theme = culte.theme or Theme.par_defaut()
        self.image_fond = lire_image(self.theme.image_fond)
        self.image_bienvenue = lire_image(self.theme.image_bienvenue)
        self.presentation = Presentation()
        self.presentation.slide_width = LARGEUR
        self.presentation.slide_height = HAUTEUR
        self.presentation.core_properties.title = culte.titre

    def exporter(self):
        for groupe in d.groupes_du_culte(self.culte):
            nom = groupe.element.nom if groupe.element else "Bienvenue"
            for diapo in groupe.diapos:
                self.ajouter(diapo, nom)
        sortie = io.BytesIO()
        self.presentation.save(sortie)
        return sortie.getvalue()

    # -------------------------------------------------------------- Diapos

    def ajouter(self, diapo, nom):
        diapositive = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[MISE_EN_PAGE_VIERGE]
        )
        self.fond(diapositive)

        if diapo.type == d.BIENVENUE and self.image_bienvenue:
            self.image_pleine_page(diapositive, self.image_bienvenue)
        elif diapo.type == d.BIENVENUE:
            self.bandeau(diapositive)
            self.texte(diapositive, diapo, italique=True, police=self.theme.police_titres, gras=True)
        elif diapo.type == d.TITRE:
            self.texte(diapositive, diapo, police=self.theme.police_titres, gras=self.theme.gras)
        else:
            self.texte(diapositive, diapo, police=self.theme.police, gras=self.theme.gras)

        if diapo.masque:
            # Diapo masquée de PowerPoint : dans le fichier, pas en projection.
            diapositive._element.set("show", "0")
        notes = f"{nom} — {diapo.libelle}" if diapo.libelle and diapo.libelle != nom else nom
        diapositive.notes_slide.notes_text_frame.text = notes

    def fond(self, diapositive):
        remplissage = diapositive.background.fill
        remplissage.solid()
        remplissage.fore_color.rgb = rgb(self.theme.couleur_fond)
        if self.image_fond:
            self.image_pleine_page(diapositive, self.image_fond)

    def image_pleine_page(self, diapositive, contenu):
        diapositive.shapes.add_picture(io.BytesIO(contenu), 0, 0, LARGEUR, HAUTEUR)

    def bandeau(self, diapositive):
        """Bandeau gris encadré de deux filets, comme la diapo de bienvenue actuelle."""
        gauche = pt((rendu.LARGEUR_PT - rendu.LARGEUR_BANDEAU_PT) / 2)
        haut = pt((rendu.HAUTEUR_PT - rendu.HAUTEUR_BANDEAU_PT) / 2)
        largeur, hauteur = pt(rendu.LARGEUR_BANDEAU_PT), pt(rendu.HAUTEUR_BANDEAU_PT)
        filet = pt(4)
        formes = diapositive.shapes
        for y, h, couleur in (
            (haut - filet * 2, filet, "#c8c8c8"),
            (haut, hauteur, self.theme.couleur_bandeau),
            (haut + hauteur + filet, filet, "#c8c8c8"),
        ):
            forme = formes.add_shape(MSO_SHAPE.RECTANGLE, gauche, y, largeur, h)
            forme.fill.solid()
            forme.fill.fore_color.rgb = rgb(couleur)
            forme.line.fill.background()
            forme.shadow.inherit = False

    def texte(self, diapositive, diapo, police, gras, italique=False):
        largeur, hauteur = rendu.zone_texte(diapo)
        zone = diapositive.shapes.add_textbox(
            pt((rendu.LARGEUR_PT - largeur) / 2), pt((rendu.HAUTEUR_PT - hauteur) / 2), pt(largeur), pt(hauteur)
        )
        cadre = zone.text_frame
        cadre.word_wrap = True
        cadre.vertical_anchor = MSO_ANCHOR.MIDDLE
        cadre.margin_left = cadre.margin_right = cadre.margin_top = cadre.margin_bottom = 0

        taille = Pt(round(rendu.taille_pt(diapo, self.theme)))
        majuscules = rendu.en_majuscules(diapo, self.theme)
        for i, ligne in enumerate(diapo.lignes or [""]):
            paragraphe = cadre.paragraphs[0] if i == 0 else cadre.add_paragraph()
            paragraphe.alignment = PP_ALIGN.CENTER
            paragraphe.line_spacing = 1.0
            morceau = paragraphe.add_run()
            morceau.text = ligne.upper() if majuscules else ligne
            fonte = morceau.font
            fonte.name = police
            fonte.size = taille
            fonte.bold = gras
            fonte.italic = italique
            fonte.color.rgb = rgb(self.theme.couleur_texte)


def exporter(culte):
    """Contenu binaire du fichier .pptx du culte."""
    return Exporteur(culte).exporter()


def nom_de_fichier(culte):
    interdits = '\\/:*?"<>|'
    nom = "".join("-" if c in interdits else c for c in culte.titre).strip() or "culte"
    return f"{nom}.pptx"
